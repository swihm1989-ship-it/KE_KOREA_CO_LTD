"""Read supplied workbooks and prepare a private, reproducible PostgreSQL seed.

No formulas, document instructions, network requests, or database writes execute.
"""
import argparse
from collections import Counter
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import uuid

NAMESPACE = uuid.UUID("7d271b5d-e68f-4dfa-919c-b18a797a2377")
TOLERANCE = Decimal("0.000001")


def text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split()) or None


def number(value, label):
    if value is None or isinstance(value, bool):
        raise ValueError(f"Missing or invalid {label}")
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"Missing or invalid {label}") from exc
    if not result.is_finite() or result < 0:
        raise ValueError(f"Missing or invalid {label}")
    if result != result.quantize(TOLERANCE):
        raise ValueError(f"Unsupported precision in {label}")
    return result


def decimal_text(value):
    return format(value, "f").rstrip("0").rstrip(".") if "." in format(value, "f") else format(value, "f")


def identity(code, name, capacity, kind):
    if not text(code) or not text(name):
        raise ValueError("Missing product code or item name")
    key = json.dumps(["ECLADO", kind, *[(text(v) or "").casefold()
                                     for v in (code, name, capacity)]], ensure_ascii=False)
    return str(uuid.uuid5(NAMESPACE, key)), hashlib.sha256(key.encode()).hexdigest()


def build_catalog(purchase_rows, buyer_rows, purchase_samples, buyer_samples):
    if len(purchase_rows) != len(buyer_rows) or len(purchase_samples) != len(buyer_samples):
        raise ValueError("Source row counts do not match")
    products, purchases, sales, issues, seen = [], [], [], [], set()
    group = None

    def add(code, name, capacity, kind, purchase, sale, reference, row_number, source_reference):
        product_id, key = identity(code, name, capacity, kind)
        if product_id in seen:
            raise ValueError(f"Duplicate product identity at {kind} row {row_number}")
        seen.add(product_id)
        cost = number(purchase, "purchase price")
        base = number(sale, "sale price")
        discount = Decimal("10") if kind == "standard" else None
        final = base * Decimal("0.9") if discount is not None else None
        if discount is not None:
            try:
                cached = Decimal(str(reference))
            except InvalidOperation as exc:
                raise ValueError(f"Missing reference discount at row {row_number}") from exc
            if not cached.is_finite() or abs(cached - final) > TOLERANCE:
                raise ValueError(f"Reference discount mismatch at row {row_number}")
        products.append({"id": product_id, "identity_key": key, "brand": "ECLADO",
                         "kind": kind, "source_code": text(code), "item_name": text(name),
                         "capacity": text(capacity), "group_name": group if kind == "standard" else "Samples",
                         "metadata": {"document_language": "en", "capacity_status":
                                      "provided" if text(capacity) else "not_provided"}})
        common = {"product_id": product_id, "source_sheet": "ECLADO_Products_List" if kind == "standard" else "Sample",
                  "source_row": row_number}
        purchases.append({**common, "supplier_key": "ECLADO", "currency": "KRW",
                          "unit_price": decimal_text(cost), "source_column": "N" if kind == "standard" else "C",
                          "source_reference": source_reference})
        sales.append({**common, "buyer_key": "DEAR_GLOW", "currency": "USD" if kind == "standard" else "KRW",
                      "unit_price": decimal_text(base), "source_column": "N" if kind == "standard" else "C",
                      "reference_discount_percent": decimal_text(discount) if discount is not None else None,
                      "reference_discounted_price": decimal_text(final) if final is not None else None,
                      "source_reference": {"role": "buyer_workbook", "purchase_workbook_cross_check_column":
                                           "P" if kind == "standard" else None,
                                           "discount_is_reference_only": True}})

    for n, (purchase, buyer) in enumerate(zip(purchase_rows, buyer_rows), 4):
        if len(purchase) < 17 or len(buyer) < 15:
            raise ValueError(f"Unexpected source layout at row {n}")
        if [text(v) for v in purchase[1:4]] != [text(v) for v in buyer[1:4]]:
            raise ValueError(f"Source product identity mismatch at row {n}")
        if text(purchase[0]):
            group = text(purchase[0])
        if number(purchase[15], "sale price") != number(buyer[13], "sale price"):
            raise ValueError(f"Source sale price mismatch at row {n}")
        # Check both discount columns; never evaluate their Excel formulas.
        for reference in (purchase[16], buyer[14]):
            try:
                cached = Decimal(str(reference))
            except InvalidOperation as exc:
                raise ValueError(f"Missing reference discount at row {n}") from exc
            if not cached.is_finite() or abs(cached - number(buyer[13], "sale price") * Decimal(".9")) > TOLERANCE:
                raise ValueError(f"Reference discount mismatch at row {n}")
        packing = dict(zip(("unit_weight_kg", "units_per_carton", "carton_net_weight_kg",
                            "carton_gross_weight_kg", "carton_dimensions_mm", "manufacturer", "ean13"), purchase[4:11]))
        invalid_packing = packing["carton_dimensions_mm"] is not None and not isinstance(packing["carton_dimensions_mm"], str)
        if invalid_packing:
            issues.append({"type": "invalid_packing_reference", "source_sheet": "ECLADO_Products_List",
                           "source_row": n, "code": text(purchase[1]),
                           "action": "Confirm actual carton dimensions and weights before issuing a packing list."})
        add(*purchase[1:4], "standard", purchase[13], buyer[13], buyer[14], n,
            {"role": "purchase_workbook", "packing_reference": packing,
             "packing_status": "needs_review" if invalid_packing else "reference_only_not_actual_shipment"})

    for n, (purchase, buyer) in enumerate(zip(purchase_samples, buyer_samples), 2):
        if len(purchase) < 3 or len(buyer) < 3 or [text(v) for v in purchase[:2]] != [text(v) for v in buyer[:2]]:
            raise ValueError(f"Source sample identity mismatch at row {n}")
        add(purchase[0], purchase[1], None, "sample", purchase[2], buyer[2], None, n,
            {"role": "purchase_workbook", "capacity_status": "not_provided"})
    for (kind, code), count in Counter((p["kind"], p["source_code"]) for p in products).items():
        if count > 1:
            issues.append({"type": "duplicate_code", "kind": kind, "code": code,
                           "count": count, "action": "Select by product ID, item name and capacity; do not merge by code."})
    if purchase_samples:
        issues.append({"type": "missing_sample_capacity", "count": len(purchase_samples),
                       "action": "Keep sample capacity null until confirmed; do not reuse full-size capacity."})
    return {"products": products, "purchase_prices": purchases, "buyer_prices": sales, "issues": issues}


def make_seed(catalog, sources):
    fingerprint = json.dumps(sources, sort_keys=True, ensure_ascii=False)
    import_id = str(uuid.uuid5(NAMESPACE, fingerprint))
    payload = json.loads(json.dumps(catalog, ensure_ascii=False))
    payload_sha256 = hashlib.sha256(json.dumps(
        {"catalog": catalog, "sources": sources}, ensure_ascii=False, sort_keys=True
    ).encode("utf-8")).hexdigest()
    payload["import"] = {"id": import_id, "source_label": "260511 (file-name label; not an effective date)",
                         "source_files": sources, "validation_issues": catalog["issues"],
                         "payload_sha256": payload_sha256}
    for kind in ("purchase_prices", "buyer_prices"):
        for price in payload[kind]:
            price["import_id"] = import_id
            price["id"] = str(uuid.uuid5(NAMESPACE, f'{import_id}:{kind}:{price["product_id"]}'))
    encoded = json.dumps(payload, ensure_ascii=False, indent=2)
    if "$kekorea_seed$" in encoded or "$kekorea_import$" in encoded:
        raise ValueError("Source contains a reserved SQL delimiter")
    template = Path(__file__).resolve().parents[1] / "supabase" / "seed_template.sql"
    sql = template.read_text(encoding="utf-8").replace("__CATALOG_JSON__", encoded)
    report = {"status": "prepared_not_applied", "import_id": import_id, "source_files": sources,
              "payload_sha256": payload_sha256,
              "counts": {"products": len(payload["products"]),
                         "standard_products": sum(p["kind"] == "standard" for p in payload["products"]),
                         "sample_products": sum(p["kind"] == "sample" for p in payload["products"]),
                         "purchase_prices": len(payload["purchase_prices"]), "buyer_prices": len(payload["buyer_prices"])},
              "issues": catalog["issues"]}
    return sql, report


def validate_price_header(value, role):
    header = (text(value) or "").upper()
    if role == "purchase":
        valid = "PRICE" in header and any(mark in header for mark in ("KRW", "\\", "₩"))
        valid = valid and "USD" not in header and "$" not in header
        currency = "KRW"
    elif role == "buyer":
        valid = "PRICE" in header and "USD" in header and "KRW" not in header
        currency = "USD"
    else:
        valid, currency = False, None
    if not valid:
        raise ValueError("Unexpected primary price header/currency")
    return currency


def validate_counts(catalog, standard_count, sample_count):
    counts = Counter(p["kind"] for p in catalog["products"])
    if counts["standard"] != standard_count or counts["sample"] != sample_count:
        raise ValueError(f"Source product count mismatch: expected {standard_count} standard / {sample_count} sample; got {dict(counts)}")


def read_workbook(path, role):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        main = workbook["ECLADO_Products_List"]
        if text(main["C3"].value) != "Items" or text(main["D3"].value) != "Capacity":
            raise ValueError("Unexpected item/capacity headers")
        validate_price_header(main["N3"].value, role)
        if role == "purchase" and "DEAR GLOW" not in (text(main["P3"].value) or "").upper():
            raise ValueError("Missing Dear Glow price cross-check column")
        samples = workbook["Sample"]
        if "KRW" not in (text(samples["C1"].value) or "").upper():
            raise ValueError("Unexpected sample price currency")
        # Keep physical row offsets intact rather than silently dropping malformed rows.
        rows = list(main.iter_rows(min_row=4, values_only=True))
        sample_rows = list(samples.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    source = {"role": role, "file_name": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
    return rows, sample_rows, source


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--purchase", required=True, type=Path)
    parser.add_argument("--buyer", required=True, type=Path)
    parser.add_argument("--expected-standard-count", type=int, default=131)
    parser.add_argument("--expected-sample-count", type=int, default=26)
    parser.add_argument("--out-dir", type=Path, default=Path(__file__).resolve().parents[1] / "private" / "eclado")
    args = parser.parse_args()
    purchase, purchase_samples, purchase_source = read_workbook(args.purchase, "purchase")
    buyer, buyer_samples, buyer_source = read_workbook(args.buyer, "buyer")
    catalog = build_catalog(purchase, buyer, purchase_samples, buyer_samples)
    validate_counts(catalog, args.expected_standard_count, args.expected_sample_count)
    sql, report = make_seed(catalog, [purchase_source, buyer_source])
    args.out_dir.mkdir(parents=True, exist_ok=True)
    (args.out_dir / "002_eclado_catalog_seed.sql").write_text(sql, encoding="utf-8")
    (args.out_dir / "import_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": report["status"], "import_id": report["import_id"],
                      "counts": report["counts"], "issue_types": [i["type"] for i in report["issues"]]}, indent=2))


if __name__ == "__main__":
    main()
